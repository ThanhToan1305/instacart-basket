"""Preserve the reference cover/styles; write a concise, evidence-based report."""
import json,csv,hashlib,sys,os,subprocess,shutil
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]
REFERENCE=ROOT/'docs/reference/Bao_cao_prototype_Big_Data_Instacart.docx'
TARGET=ROOT/'docs/Bao_cao_prototype_hoan_thien.docx'

def read(name):return json.loads((ROOT/'outputs/metrics'/name).read_text())

def main():
    original_hash=hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    d=Document(REFERENCE);body=d._element.body
    cutoff=next(p._p for p in d.paragraphs if p.text=='MỤC LỤC')
    index=list(body).index(cutoff)
    for element in list(body)[index:]:
        if element.tag!=qn('w:sectPr'):body.remove(element)
    for p in d.paragraphs:
        if p.text.startswith('Phiên bản 0.1'):
            p.text='Phiên bản hoàn thiện — Kết quả thực nghiệm và bằng chứng local Spark'
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.1
    for style in ['Heading 1','Heading 2','Heading 3']:
        d.styles[style].paragraph_format.space_before=Pt(7);d.styles[style].paragraph_format.space_after=Pt(5)
        d.styles[style].paragraph_format.page_break_before=False
    settings=d.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    for section_index,section in enumerate(d.sections):
        section.page_width=Cm(21);section.page_height=Cm(29.7)
        section.left_margin=Cm(2.7);section.right_margin=Cm(2);section.top_margin=Cm(2);section.bottom_margin=Cm(2)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for child in list(footer._p):
            if child.tag!=qn('w:pPr'):footer._p.remove(child)
        footer.add_run('Trang ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
        page_number=section._sectPr.find(qn('w:pgNumType'))
        if page_number is None:page_number=OxmlElement('w:pgNumType');section._sectPr.append(page_number)
        page_number.set(qn('w:start'),str(section_index+1))
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                run.text=run.text.replace('PROTOTYPE v0.1','THỰC NGHIỆM FULL').replace('Prototype v0.1','Thực nghiệm full')
    raw=read('raw_profile.json');eda=read('eda_summary.json');model=read('model_summary.json');prototype=read('prototype_summary.json');env=read('environment.json');tests=read('test_summary.json');sample=read('sample_integration.json');k=eda['kpi'];c=model['selected']
    figure_number=0;table_number=0
    def para(text):d.add_paragraph(text)
    def new_page():d.add_page_break()
    def heading(text,level=1):d.add_heading(text,level)
    def table(headers,rows,caption,source):
        nonlocal table_number
        table_number+=1;para(f'Bảng {table_number}. {caption}')
        t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for row in t.rows:
            trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                    for r in p.runs:r.font.size=Pt(9)
        widths=[16.3/len(headers)]*len(headers)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        p=d.add_paragraph('Nguồn: '+source);p.paragraph_format.space_after=Pt(4)
        for r in p.runs:r.italic=True;r.font.size=Pt(9)
    def image(relative,caption,max_height=7):
        nonlocal figure_number
        path=ROOT/relative
        with Image.open(path) as im:width,height=im.size
        width_cm=min(16.0,max_height*width/height)
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(str(path),width=Cm(width_cm))
        figure_number+=1
        p=d.add_paragraph(f'Hình {figure_number}. {caption}. Nguồn: {relative}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:r.font.size=Pt(9)
    # Page 2: automatic field with a readable cached placeholder; Word refreshes Heading 1–3.
    d.add_paragraph('MỤC LỤC',style='Title')
    p=d.add_paragraph();run=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-3" \\h \\z \\u ';run._r.append(instr)
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate');run._r.append(separate)
    entries=[('MỞ ĐẦU VÀ CHƯƠNG 1',3),('CHƯƠNG 2. DỮ LIỆU VÀ KIẾN TRÚC',4),('CHƯƠNG 3. PHƯƠNG PHÁP',5),('CHƯƠNG 4. PROTOTYPE ĐÃ TRIỂN KHAI',6),('CHƯƠNG 5. THỰC NGHIỆM',7),('5.1. Môi trường thực nghiệm',7),('5.2. Dataset và tiền xử lý',8),('5.3. Kết quả EDA và KPI',9),('5.4. FP-Growth và lựa chọn tham số',11),('5.5. Đánh giá luật và gợi ý',12),('5.6. Dashboard prototype',13),('5.7. Hiệu năng và kiểm thử',14),('5.8. Hạn chế và hướng phát triển',15),('KẾT LUẬN VÀ TÀI LIỆU THAM KHẢO',15)]
    for text,page in entries:d.add_paragraph(f'{text}\t{page}')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');d.paragraphs[-1].add_run()._r.append(end)
    para('Mục lục là trường TOC liên kết Heading 1–3. Cho phép cập nhật trường khi mở Word; có thể Ctrl+A, F9 để làm mới số trang.')
    new_page();heading('MỞ ĐẦU')
    para('Đề tài khai phá hành vi mua sắm và gợi ý sản phẩm mua kèm trên Instacart. Mục tiêu là tạo pipeline batch có thể tái lập, kiểm tra chất lượng, sinh luật giải thích được và trình bày trên dashboard. Các kết quả trong bản này là kết quả thật, truy ngược được tới Parquet, JSON, CSV và log; không sử dụng wireframe làm bằng chứng.')
    heading('CHƯƠNG 1. TỔNG QUAN BÀI TOÁN VÀ NỀN TẢNG')
    heading('1.1. Market Basket Analysis',2)
    para('Mỗi đơn hàng là một tập sản phẩm. FP-Growth tìm các itemset phổ biến mà không sinh hàng loạt ứng viên như Apriori. Luật X → Y biểu diễn quan hệ đồng xuất hiện; đây không phải kết luận nhân quả hay xác suất cá nhân đã hiệu chỉnh.')
    table(['Chỉ số','Công thức','Vai trò'],[['Support','P(X ∪ Y)','Độ phổ biến toàn luật'],['Confidence','P(Y | X)','Độ tin cậy trên giỏ chứa X'],['Lift','Confidence / P(Y)','Liên hệ dương khi > 1']],'Chỉ số luật kết hợp','Phương pháp FP-Growth; Spark associationRules')
    heading('1.2. Spark và phạm vi triển khai',2)
    para('Spark SQL xử lý schema, chất lượng, join và tổng hợp; Spark MLlib học FP-Growth. Triển khai thực tế dùng local[*] và filesystem local với Parquet Snappy. HDFS/Spark cluster trong bản tham chiếu là kiến trúc nâng cấp, chưa được triển khai và không được coi là kết quả thực nghiệm.')
    new_page();heading('CHƯƠNG 2. DỮ LIỆU VÀ KIẾN TRÚC')
    table(['File thực tế','Dòng CSV','Dòng Bronze'],[[Path(v['source_path']).name,f"{v['row_count']:,}",f"{v['bronze_rows']:,}"] for v in raw['tables'].values()],'Dataset thực tế và đối soát','outputs/metrics/raw_profile.json')
    para('Sáu CSV liên kết qua order_id, user_id, product_id, aisle_id, department_id. Đơn test không có chi tiết sản phẩm; mã ngày 0–6 đã ẩn danh. Dữ liệu được giữ nguyên tại raw; Bronze thêm metadata, Silver chuẩn hóa và Gold phục vụ phân tích/gợi ý.')
    image('outputs/evidence/01_project_structure.png','Cấu trúc workspace thật',max_height=7)
    new_page();heading('CHƯƠNG 3. PHƯƠNG PHÁP VÀ QUY TRÌNH')
    image('outputs/evidence/02_pipeline_architecture.png','Kiến trúc local thực tế',max_height=6)
    heading('3.1. Tiền xử lý và basket',2)
    para('Khai báo schema số nguyên; kiểm tra null/duplicate/miền giá trị và khóa ngoại. Ghép orders, order products, products, aisles, departments; chỉ giữ cột cần thiết. collect_set(product_id) tạo giỏ không trùng. Chỉ giỏ prior có ít nhất hai sản phẩm được dùng học mô hình; product_name chỉ gắn khi xuất cuối.')
    heading('3.2. Gợi ý và đánh giá',2)
    para('Antecedent phải là tập con đầy đủ của giỏ. Loại sản phẩm đã có; mỗi sản phẩm dùng luật tốt nhất, xếp confidence giảm dần, rồi lift, support và mã sản phẩm. Giữ Top-5 duy nhất. Cách xếp hạng thực tế này thay cho công thức minh họa trong bản tham chiếu. Ground truth là item của đơn train, ngữ cảnh là giỏ prior gần nhất. Train cũng dùng chọn cấu hình nên là validation nội bộ.')
    new_page();heading('CHƯƠNG 4. PROTOTYPE ĐÃ TRIỂN KHAI')
    para('Đã triển khai pipeline có resume, timestamp/run_id, số dòng, log và dừng khi lỗi nghiêm trọng. Các worker FP-Growth được cô lập để một cấu hình OOM không làm mất các mô hình đã thành công. Dashboard có năm trang, cache dữ liệu Gold/metrics và không khởi động Spark.')
    table(['Thành phần','Trạng thái','Bằng chứng'],[['Raw/Bronze/Silver','Đọc được, đối soát PASS','raw_profile, silver_profile'],['KPI và EDA','Số liệu full; 12 PNG','eda_summary; outputs/figures'],['FP-Growth','6 PASS, 3 skip OOM','model_comparison.csv'],['Dashboard','5 trang; Top-5 thật','Ảnh 10–14 và 17'],['Kiểm thử','Full + mẫu raw→Gold','test_summary; sample_integration']],'Trạng thái triển khai','JSON/CSV/log thực nghiệm; Evidence Manifest')
    image('outputs/evidence/10_dashboard_overview.png','Dashboard thật, không phải wireframe',max_height=8)
    new_page();heading('CHƯƠNG 5. THỰC NGHIỆM VÀ ĐÁNH GIÁ')
    heading('5.1. Môi trường thực nghiệm',2)
    para(f"Python {env['python']}; Java {env['java'].splitlines()[0]}; PySpark {env['packages']['pyspark']}; {env['logical_cpus']} CPU logic; RAM {env['ram_total_bytes']/1024**3:.2f} GiB; đĩa trống lúc kiểm tra {env['disk_free_bytes']/1024**3:.2f} GiB. Nguồn: environment.json. Spark local[*], AQE bật, shuffle 32. FP-Growth full dùng driver 3g, 128 phân vùng, task.cpus=4; các action bổ sung dùng 1g.")
    para('execution_mode=full: tái sử dụng toàn bộ dữ liệu đã đối soát và hai mức support đã fit thành công. Một lần kiểm thử riêng dùng sample 1% theo order_id % 10000 < 100, không được gọi là full và không thay KPI full.')
    image('outputs/evidence/09_spark_ui_environment.png','Spark UI Environment thật tại localhost:4040',max_height=8)
    new_page();heading('5.2. Dataset và tiền xử lý',2)
    para(f"Full có {k['total_orders']:,} đơn, {k['total_users']:,} người dùng và {k['total_order_items']:,} lượt mua quan sát. {k['orders_without_observed_items']:,} đơn test không có item; không tạo giỏ rỗng giả. days_since_prior_order được phép null ở đơn đầu; giá trị ngày bị giới hạn 30 trong nguồn.")
    image('outputs/evidence/05_row_reconciliation.png','CSV = Bronze ở cả sáu bảng',max_height=5)
    image('outputs/evidence/06_silver_quality.png','Khóa ngoại và basket integrity',max_height=5)
    para('Bronze hiện có ingestion_time, source_file, run_id. Với dữ liệu legacy, run_id được thêm trong lần migration và ghi rõ provenance; ingestion_time cũ được giữ. Bản Bronze trước migration được giữ trong thư mục backup ẩn. Không diễn giải run_id migration thành thời điểm ingest ban đầu.')
    new_page();heading('5.3. Kết quả EDA và KPI',2)
    table(['KPI','Kết quả full'],[['Người dùng',f"{k['total_users']:,}"],['Đơn / lượt mua',f"{k['total_orders']:,} / {k['total_order_items']:,}"],['Sản phẩm / giỏ quan sát',f"{k['total_products']:,} / {k['observed_baskets']:,}"],['Giỏ trung bình / trung vị',f"{k['avg_basket_size']:.4f} / {k['median_basket_size']:.0f}"],['Reorder',f"{k['reorder_rate']:.4%}"],['Đơn/người; ngày giữa đơn',f"{k['avg_orders_per_user']:.4f}; {k['avg_days_between_orders']:.4f}"]],'KPI thực nghiệm','outputs/metrics/eda_summary.json; Gold kpi_summary')
    image('outputs/figures/01_orders_by_hour.png','Đơn hàng theo giờ',max_height=5.2)
    image('outputs/figures/05_basket_size_distribution.png','Phân phối giỏ',max_height=5.2)
    para(f"Giờ {k['peak_hour']} có {k['peak_hour_orders']:,} đơn, là đỉnh quan sát. Trung bình giỏ lớn hơn trung vị, cho thấy giỏ lớn kéo trung bình lên. Không suy luận múi giờ hoặc nguyên nhân từ dataset ẩn danh.")
    new_page();heading('5.3. Kết quả EDA và KPI (tiếp)',2)
    image('outputs/figures/03_top_20_products.png','Top sản phẩm',max_height=7)
    image('outputs/figures/07_orders_per_user_distribution.png','Số đơn mỗi người dùng',max_height=5.5)
    para(f"Banana đứng đầu với {k['most_purchased_product_count']:,} lượt mua. Aisle {k['most_purchased_aisle_name']} đứng đầu với {k['most_purchased_aisle_count']:,} lượt (top_20_aisles.csv); department phổ biến và reorder được trình bày trong hình 04/06. Phân phối số đơn/người có tổng người và tổng đơn khớp KPI; đây là action Spark mới trên orders full.")
    new_page();heading('5.4. FP-Growth và lựa chọn tham số',2)
    para(f"Có {model['preparation']['prior_baskets']:,} giỏ prior; loại {model['preparation']['excluded_single_item_baskets']:,} giỏ một sản phẩm, học trên {c['training_baskets']:,} giỏ kiểu ArrayType(IntegerType). Ba confidence chia sẻ cùng một lần fit cho mỗi support.")
    grid=[]
    for r in model['comparisons']:
        if r['status']=='PASS':grid.append([r['min_support'],r['min_confidence'],r['useful_rules'],f"{r['median_lift']:.3f}",f"{r['coverage']:.2%}",f"{r['hitrate_at_5']:.2%}"])
        else:grid.append([r['min_support'],r['min_confidence'],'SKIP OOM','—','—','—'])
    table(['Support','Conf.','Luật','Lift TV','Coverage','HitRate@5'],grid,'So sánh chín cấu hình','outputs/tables/model_comparison.csv; model_summary.json')
    para(f"Chọn support {c['min_support']}, confidence {c['min_confidence']}: {c['frequent_itemsets']:,} itemset, {c['useful_rules']:,} luật. Điểm đa tiêu chí gồm HitRate@5 (0,40), Coverage (0,35), confidence TB (0,10), lift TV chuẩn hóa (0,10), số luật hữu dụng (0,05), phạt thời gian (0,02). Công thức đầy đủ trong model_summary.json. Mô hình thắng theo cân bằng chất lượng/độ phủ/thời gian, không theo một tiêu chí số luật.")
    image('outputs/figures/08_rule_count_comparison.png','Số luật theo cấu hình PASS',max_height=5)
    para('0,0005 có bằng chứng OOM ở lần thử trước và gián đoạn ở các lần sau; không có metrics hoàn chỉnh. Các ô thiếu không phải số 0 và không chứng minh cấu hình này có chất lượng thấp.')
    new_page();heading('5.5. Đánh giá luật và gợi ý',2)
    para(f"Trên {c['eligible_users']:,} người dùng có train: Coverage={c['covered_users']:,}/{c['eligible_users']:,}={c['coverage']:.4%}; HitRate@5={c['hit_users']:,}/{c['eligible_users']:,}={c['hitrate_at_5']:.4%}. Mẫu số gồm cả người không có gợi ý. Có {c['recommendation_count']:,} gợi ý được sinh; HitRate có điều kiện {c['conditional_hitrate_at_5']:.4%} chỉ là chỉ số phụ.")
    para(f"Luật cuối lift>1, đủ confidence, không rỗng/trùng/giao hai vế, consequent một sản phẩm. Lift TB {c['avg_lift']:.4f}, trung vị {c['median_lift']:.4f}; trung bình lớn hơn trung vị do đuôi luật lift lớn. Train đồng thời dùng lựa chọn cấu hình: chưa có phép thử độc lập cuối.")
    image('outputs/figures/10_lift_distribution.png','Lift theo khoảng log₂',max_height=5)
    example=list(csv.DictReader((ROOT/'outputs/tables/recommendation_example.csv').open()))
    table(['Top-5 đơn 94','Support','Confidence','Lift'],[[r['product_name'],f"{float(r['support']):.6f}",f"{float(r['confidence']):.4f}",f"{float(r['lift']):.4f}"] for r in example],'Gợi ý từ giỏ bốn sản phẩm quan sát','Silver baskets order_id=94; recommendation_example.csv')
    para('Ngữ cảnh lấy từ giỏ thật; các gợi ý không nằm trong giỏ, khớp đủ antecedent và có luật nguồn. Lift lớn là tương quan đồng xuất hiện, không phải quan hệ nhân quả.')
    new_page();heading('5.6. Dashboard prototype',2)
    para('Dashboard có Tổng quan, Hành vi mua sắm, Luật kết hợp, Gợi ý sản phẩm và Hiệu năng hệ thống. Lọc support/confidence/lift trên mô hình Gold đã chọn, tải CSV; các đồ thị Plotly tương tác. Cache theo dấu file; thiếu dữ liệu sẽ chỉ pipeline cần chạy, không crash.')
    image('outputs/evidence/13_dashboard_recommendation.png','Dashboard thật đã chọn giỏ và hiện Top-5',max_height=10)
    para('Điểm là confidence của luật tốt nhất; tie-break lift/support/product_id. Không có luật thì fallback dựa trên top 20 sản phẩm phổ biến thực tế và ghi rõ nguồn, để trống chỉ số luật. Dashboard không chạy Spark. Các ảnh 10–14 trong manifest là ảnh Chromium thật; ảnh dài giữ đầy đủ tại outputs/evidence để phóng to đọc chi tiết.')
    new_page();heading('5.7. Hiệu năng và kiểm thử',2)
    times=[['Raw + Bronze',raw['total_seconds']],['Silver',read('silver_profile.json')['total_seconds']],['EDA/KPI',eda['total_seconds']],['Fit support=0,002',698.473],['Fit support=0,001',c['training_seconds']],['Đánh giá/xuất/biểu đồ',model['evaluation_pipeline_seconds']],['Full resume + test mới',prototype['seconds']],['Test mẫu raw→Gold',json.loads(Path(sample['summary_path']).read_text())['seconds']]]
    table(['Thành phần','Giây'],[[n,f'{s:.3f}'] for n,s in times],'Thời gian thật, không cộng lặp confidence','raw/silver/eda/model/prototype_summary và sample_integration')
    para(f"Full chạy {tests['count']} test PASS trong {tests['seconds']:.3f}s; sample riêng chạy 32 test PASS trên 34.299 đơn và 339.165 item. Mẫu chọn theo order_id, giữ đầy đủ dimension. Kiểm tra schema/null/duplicate/orphan, giỏ/luật/metrics, recommender, dashboard và fail-fast. Fragment Parquet 0 dòng hợp lệ được giữ schema và không coi là lỗi null.")
    image('outputs/evidence/07_spark_ui_jobs.png','Spark UI thật sau action full; không phải ảnh huấn luyện ban đầu',max_height=5.5)
    image('outputs/evidence/18_test_summary.png','Tổng kết test thực tế',max_height=4)
    new_page();heading('5.8. Hạn chế và hướng phát triển',2)
    for text in ['Triển khai local Spark, chưa có HDFS/cluster, deploy/auth hoặc benchmark hệ thống phân tán.',
                 'Support 0,0005 chưa hoàn tất trên full; cần môi trường đủ RAM. Không cộng các lần gián đoạn thành một thời gian end-to-end giả.',
                 'Train dùng chọn cấu hình; cần test độc lập theo thời gian/người dùng và baseline popularity, Precision/Recall/NDCG@5.',
                 'Confidence chưa hiệu chỉnh thành xác suất cá nhân; nhiệm vụ loại item đã có không bao quát toàn bộ nhu cầu reorder.',
                 'Fallback top 20 có thể ít hơn năm khi nhiều item đã trong giỏ; cần popularity rộng hơn và kiểm tra tồn kho.',
                 'Resume fingerprint dùng tên/kích thước/mtime raw và hash config; hướng production cần hash dataset và lineage đầy đủ.']:
        d.add_paragraph(text,style='List Bullet')
    heading('KẾT LUẬN')
    para('Prototype đã có pipeline được đối soát, KPI/luật thật, dashboard giải thích được, kiểm thử full và mẫu độc lập, ảnh Spark UI/dashboard thật. Các số liệu full trong chương khác với sample kiểm thử. Evidence Manifest ghi run_id, thời gian, nguồn và trạng thái cho từng ảnh; báo cáo nghiệm thu ghi PASS/FAIL của từng tiêu chí. File tham chiếu giữ nguyên.')
    heading('TÀI LIỆU THAM KHẢO')
    refs=['Instacart. The Instacart Online Grocery Shopping Dataset 2017; Kaggle competition data.',
          'Han, J., Pei, J., Yin, Y. Mining Frequent Patterns without Candidate Generation. SIGMOD, 2000.',
          'Apache Spark. Frequent Pattern Mining — FP-Growth; Spark MLlib documentation.',
          'Zaharia et al. Apache Spark: A Unified Engine for Big Data Processing. CACM, 2016.',
          'Apache Hadoop. HDFS Architecture Guide (kiến trúc nâng cấp).',
          'Streamlit. Build data apps in Python.']
    for i,text in enumerate(refs,1):para(f'[{i}] {text}')
    d.save(TARGET)
    reopened=Document(TARGET)
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==original_hash
    assert len(reopened.inline_shapes)>=10
    for shape in reopened.inline_shapes:assert shape.width<=Cm(16.3) and shape.height<=Cm(24)
    assert 'TOC' in reopened._element.xml and any('PAGE' in s.footer._element.xml for s in reopened.sections)
    result={'status':'PASS','file':str(TARGET),'reference_sha256':original_hash,'images':len(reopened.inline_shapes),
            'tables':len(reopened.tables),'manual_page_breaks':len(reopened._element.xpath('.//w:br[@w:type="page"]')),
            'toc':'automatic Heading 1-3 field, update on open','page_numbers':'PAGE footer field','reference_unchanged':True}
    (ROOT/'outputs/metrics/word_validation.json').write_text(json.dumps(result,indent=2)+'\n')
    print('Word package reopened and structural checks PASS:',TARGET)

if __name__=='__main__':main()
